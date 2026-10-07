"""Génération de rapports PDF/CSV (fpdf2) — rapport de conformité flotte, coûts CSV, fiche véhicule."""
import csv
import io
from datetime import datetime, timezone

from fpdf import FPDF
from fpdf.fonts import FontFace

_BOLD = FontFace(emphasis="BOLD")

LEVELS_FR = {"ok": "Conforme", "warning": "Avertissement", "critical": "Critique",
             "expired": "Expiré", "unknown": "Incomplet"}

_REPL = {"\u2014": "-", "\u2013": "-", "\u2019": "'", "\u00a0": " ", "\u202f": " ",
         "\u2026": "...", "\u20ac": "EUR"}


def _tx(s):
    s = "" if s is None else str(s)
    for a, b in _REPL.items():
        s = s.replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


def _date_fr(iso):
    if not iso:
        return "-"
    try:
        return datetime.fromisoformat(str(iso)[:10]).strftime("%d.%m.%Y")
    except ValueError:
        return str(iso)


def _due(iso, days):
    if not iso:
        return "-"
    txt = _date_fr(iso)
    if days is None:
        return txt
    return f"{txt} (échu)" if days < 0 else f"{txt} ({days} j)"


def _chf(x):
    try:
        n = float(x or 0)
    except (TypeError, ValueError):
        n = 0
    if not n:
        return "-"
    return f"{n:,.0f}".replace(",", "'") + " CHF"


class _Report(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("helvetica", "I", 8)
        self.set_text_color(130)
        self.cell(0, 8, _tx(f"LogiTrak Documents — Rapport de conformité — page {self.page_no()}"), align="C")
        self.set_text_color(0)


def _section(pdf, title):
    pdf.set_font("helvetica", "B", 12)
    pdf.cell(0, 8, _tx(title), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(1)


def build_conformity_pdf(vehicles: list) -> bytes:
    pdf = _Report(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()

    pdf.set_font("helvetica", "B", 16)
    pdf.cell(0, 9, _tx("Rapport de conformité — Flotte"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 9)
    pdf.set_text_color(100)
    pdf.cell(0, 5, _tx(f"Généré le {datetime.now(timezone.utc).strftime('%d.%m.%Y %H:%M')} UTC · "
                       f"{len(vehicles)} véhicule(s) · Échéances leasing / assurance / contrôle technique et coûts"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)

    counts = {"ok": 0, "warning": 0, "critical": 0, "expired": 0, "unknown": 0}
    for v in vehicles:
        overall = (v.get("metrics") or {}).get("overall") or "unknown"
        counts[overall if overall in counts else "unknown"] += 1
    _section(pdf, "Synthèse")
    pdf.set_font("helvetica", "", 10)
    pdf.cell(0, 6, _tx(f"Conformes : {counts['ok']}   ·   Avertissements : {counts['warning']}   ·   "
                       f"Critiques : {counts['critical']}   ·   Expirés : {counts['expired']}   ·   "
                       f"Données incomplètes : {counts['unknown']}"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    _section(pdf, "Échéances & conformité")
    pdf.set_font("helvetica", "", 8.5)
    with pdf.table(col_widths=(20, 34, 26, 26, 26, 18), text_align=("LEFT", "LEFT", "LEFT", "LEFT", "LEFT", "CENTER"),
                   line_height=6, padding=1.2) as table:
        head = table.row()
        for h in ("Plaque", "Véhicule", "Fin de leasing", "Échéance assurance", "Prochain contrôle", "Statut"):
            head.cell(_tx(h))
        for v in vehicles:
            m = v.get("metrics") or {}
            leasing, assurance, controle = v.get("leasing") or {}, v.get("assurance") or {}, v.get("controle_technique") or {}
            row = table.row()
            row.cell(_tx(v.get("plaque") or "-"))
            row.cell(_tx(" ".join(x for x in (v.get("marque"), v.get("modele")) if x) or "-"))
            row.cell(_tx(_due(leasing.get("date_fin"), (m.get("leasing") or {}).get("days_remaining"))))
            row.cell(_tx(_due(assurance.get("date_echeance"), (m.get("assurance") or {}).get("days_remaining"))))
            row.cell(_tx(_due(controle.get("date_prochain"), (m.get("controle") or {}).get("days_remaining"))))
            row.cell(_tx(LEVELS_FR.get(m.get("overall"), "-")))
    pdf.ln(6)

    _section(pdf, "Coûts leasing & assurance")
    pdf.set_font("helvetica", "", 8.5)
    tot_mens, tot_rest, tot_prime = 0.0, 0.0, 0.0
    with pdf.table(col_widths=(20, 34, 24, 20, 26, 26), text_align=("LEFT", "LEFT", "RIGHT", "CENTER", "RIGHT", "RIGHT"),
                   line_height=6, padding=1.2) as table:
        head = table.row()
        for h in ("Plaque", "Véhicule", "Mensualité leasing", "Mois restants", "Coût restant leasing", "Prime assurance / an"):
            head.cell(_tx(h))
        for v in vehicles:
            m = v.get("metrics") or {}
            ml = m.get("leasing") or {}
            leasing, assurance = v.get("leasing") or {}, v.get("assurance") or {}
            mens = leasing.get("mensualite_chf") or leasing.get("cout_mensuel") or 0
            rest = ml.get("cost_remaining") or 0
            prime = assurance.get("prime_annuelle") or 0
            tot_mens += float(mens or 0)
            tot_rest += float(rest or 0)
            tot_prime += float(prime or 0)
            row = table.row()
            row.cell(_tx(v.get("plaque") or "-"))
            row.cell(_tx(" ".join(x for x in (v.get("marque"), v.get("modele")) if x) or "-"))
            row.cell(_tx(_chf(mens)))
            row.cell(_tx(ml.get("months_remaining") if ml.get("months_remaining") is not None else "-"))
            row.cell(_tx(_chf(rest)))
            row.cell(_tx(_chf(prime)))
        total = table.row()
        pdf.set_font("helvetica", "B", 8.5)
        total.cell(_tx("TOTAL"))
        total.cell(_tx(f"{len(vehicles)} véhicule(s)"))
        total.cell(_tx(_chf(tot_mens)))
        total.cell("-")
        total.cell(_tx(_chf(tot_rest)))
        total.cell(_tx(_chf(tot_prime)))
    pdf.ln(4)
    pdf.set_font("helvetica", "I", 8)
    pdf.set_text_color(110)
    pdf.cell(0, 5, _tx("Coût restant leasing = mois restants × mensualité. Les montants proviennent des contrats "
                       "saisis dans les fiches véhicules ; les lignes sans contrat affichent « - »."),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(6)

    _section(pdf, "Environnement — consommation & CO2 officiels")
    co2_vals, conso_vals = [], []
    for v in vehicles:
        if v.get("co2_g_km") not in (None, ""):
            co2_vals.append(float(v["co2_g_km"]))
        if v.get("conso_officielle_l_100km") not in (None, ""):
            conso_vals.append(float(v["conso_officielle_l_100km"]))
    pdf.set_font("helvetica", "", 10)
    if co2_vals or conso_vals:
        parts = []
        if co2_vals:
            parts.append(f"CO2 officiel moyen : {sum(co2_vals) / len(co2_vals):.0f} g/km ({len(co2_vals)}/{len(vehicles)} véhicules renseignés)")
        if conso_vals:
            parts.append(f"Conso officielle moyenne : {sum(conso_vals) / len(conso_vals):.1f} L/100 km ({len(conso_vals)}/{len(vehicles)})")
        pdf.cell(0, 6, _tx("   ·   ".join(parts)), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
    pdf.set_font("helvetica", "", 8.5)
    with pdf.table(col_widths=(20, 32, 24, 28, 24, 24), text_align=("LEFT", "LEFT", "LEFT", "LEFT", "LEFT", "LEFT"),
                   line_height=6, padding=1.2) as table:
        head = table.row()
        for h in ("Plaque", "Véhicule", "Carburant", "Conso officielle", "Conso réelle", "CO2 officiel"):
            head.cell(_tx(h))
        for v in vehicles:
            conso_off = (f"{v['conso_officielle_l_100km']} L/100 km ({v.get('conso_officielle_norme') or '-'})"
                         if v.get("conso_officielle_l_100km") not in (None, "") else "-")
            reelle = (f"{v['conso_reelle_l_100km']} L/100 km"
                      if v.get("conso_reelle_l_100km") not in (None, "") else "-")
            co2 = (f"{v['co2_g_km']} g/km ({v.get('co2_norme') or '-'})"
                   if v.get("co2_g_km") not in (None, "") else "-")
            row = table.row()
            row.cell(_tx(v.get("plaque") or "-"))
            row.cell(_tx(" ".join(x for x in (v.get("marque"), v.get("modele")) if x) or "-"))
            row.cell(_tx(v.get("type_carburant") or "-"))
            row.cell(_tx(conso_off))
            row.cell(_tx(reelle))
            row.cell(_tx(co2))
    pdf.ln(4)
    pdf.set_font("helvetica", "I", 8)
    pdf.set_text_color(110)
    pdf.cell(0, 5, _tx("Valeurs officielles issues de la base ASTRA/OFROU (copie locale) ou de la carte grise ; "
                       "conso réelle mesurée (CAN/OBD ou pleins). Les véhicules sans données affichent « - »."),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)

    return bytes(pdf.output())


def _n(x):
    try:
        f = float(x or 0)
    except (TypeError, ValueError):
        return ""
    if not f:
        return ""
    return int(f) if f == int(f) else round(f, 2)


def build_costs_csv(vehicles: list, engine_totals: dict = None) -> str:
    """CSV des coûts flotte — séparateur ';' (Excel FR/CH), dates JJ.MM.AAAA."""
    engine_totals = engine_totals or {}
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(["Plaque", "Marque", "Modèle", "Base", "Groupe", "Statut",
                "Société leasing", "Fin de leasing", "Mensualité CHF", "Mois restants", "Coût restant CHF",
                "Compagnie assurance", "Échéance assurance", "Prime annuelle CHF",
                "Prochain contrôle", "Coût annuel tous postes CHF"])
    tot_mens = tot_rest = tot_prime = tot_engine = 0.0
    for v in vehicles:
        m = v.get("metrics") or {}
        ml = m.get("leasing") or {}
        leasing, assurance = v.get("leasing") or {}, v.get("assurance") or {}
        controle = v.get("controle_technique") or {}
        mens = leasing.get("mensualite_chf") or leasing.get("cout_mensuel") or 0
        rest = ml.get("cost_remaining") or 0
        prime = assurance.get("prime_annuelle") or 0
        engine = engine_totals.get(v.get("id")) or 0
        tot_mens += float(mens or 0)
        tot_rest += float(rest or 0)
        tot_prime += float(prime or 0)
        tot_engine += float(engine or 0)
        w.writerow([
            v.get("plaque") or "", v.get("marque") or "", v.get("modele") or "",
            v.get("base") or "", v.get("groupe") or "", LEVELS_FR.get(m.get("overall"), ""),
            leasing.get("societe") or "",
            _date_fr(leasing.get("date_fin")) if leasing.get("date_fin") else "",
            _n(mens), ml.get("months_remaining") if ml.get("months_remaining") is not None else "", _n(rest),
            assurance.get("compagnie") or "",
            _date_fr(assurance.get("date_echeance")) if assurance.get("date_echeance") else "",
            _n(prime),
            _date_fr(controle.get("date_prochain")) if controle.get("date_prochain") else "",
            _n(engine),
        ])
    w.writerow(["TOTAL", "", "", "", "", "", "", "", _n(tot_mens), "", _n(tot_rest), "", "", _n(tot_prime), "", _n(tot_engine)])
    return buf.getvalue()


def _dt_fr(iso):
    try:
        return datetime.fromisoformat(str(iso).replace("Z", "+00:00")).strftime("%d.%m.%Y %H:%M")
    except (ValueError, TypeError):
        return str(iso or "-")


def _size_h(n):
    try:
        f = float(n or 0)
    except (TypeError, ValueError):
        return "-"
    if f <= 0:
        return "-"
    return f"{f / 1024:.0f} Ko" if f < 1024 * 1024 else f"{f / (1024 * 1024):.1f} Mo"


def _kv_table(pdf, pairs):
    pdf.set_font("helvetica", "", 8.5)
    with pdf.table(col_widths=(30, 70), first_row_as_headings=False,
                   borders_layout="INTERNAL", line_height=6, padding=1.2,
                   text_align=("LEFT", "LEFT")) as table:
        for label, value in pairs:
            row = table.row()
            row.cell(_tx(label), style=_BOLD)
            row.cell(_tx(value if value not in (None, "") else "-"))


def build_vehicle_pdf(vehicle: dict, history: list, documents: list, doc_type_labels: dict = None) -> bytes:
    """Fiche individuelle : identité, moteur, contrats, documents et historique."""
    m = vehicle.get("metrics") or {}
    leasing, assurance = vehicle.get("leasing") or {}, vehicle.get("assurance") or {}
    cg, controle = vehicle.get("carte_grise") or {}, vehicle.get("controle_technique") or {}
    labels = doc_type_labels or {}

    pdf = _Report(orientation="P", format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()

    pdf.set_font("helvetica", "B", 16)
    pdf.cell(0, 9, _tx(f"Fiche véhicule — {vehicle.get('plaque') or '-'}"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 9)
    pdf.set_text_color(100)
    sub = " ".join(x for x in (vehicle.get("marque"), vehicle.get("modele")) if x) or "-"
    if vehicle.get("annee"):
        sub += f" · {vehicle['annee']}"
    pdf.cell(0, 5, _tx(f"{sub} · Statut global : {LEVELS_FR.get(m.get('overall'), '-')} · "
                       f"Généré le {datetime.now(timezone.utc).strftime('%d.%m.%Y %H:%M')} UTC"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)

    _section(pdf, "Identité & exploitation")
    km = vehicle.get("kilometrage") or 0
    _kv_table(pdf, [
        ("VIN", vehicle.get("vin")),
        ("Kilométrage", f"{km:,}".replace(",", "'") + " km" if km else "-"),
        ("Base / Groupe", f"{vehicle.get('base') or '-'} / {vehicle.get('groupe') or '-'}"),
        ("Responsable", vehicle.get("responsable")),
        ("Catégorie / Variante", f"{vehicle.get('categorie') or '-'} / {vehicle.get('variante') or '-'}"),
        ("N° homologation", vehicle.get("numero_homologation")),
        ("1re mise en circulation", _date_fr(cg.get("date_mise_circulation")) if cg.get("date_mise_circulation") else "-"),
        ("Poids à vide / total", f"{vehicle.get('poids_vide') or '-'} kg / {cg.get('poids_total') or '-'} kg"),
    ])
    pdf.ln(4)

    _section(pdf, "Moteur & consommation")
    cyl = vehicle.get("cylindree_cm3")
    reelle = vehicle.get("conso_reelle_l_100km")
    sources = {"can": "CAN", "fms": "FMS", "obd": "OBD", "fuel_transactions": "Pleins carburant", "manual": "Manuelle"}
    _kv_table(pdf, [
        ("Carburant", vehicle.get("type_carburant")),
        ("Cylindrée", f"{cyl} cm3 - {cyl / 1000:.1f} L" if cyl else "-"),
        ("Puissance", f"{vehicle.get('puissance_kw')} kW" if vehicle.get("puissance_kw") else "-"),
        ("CO2 officiel", f"{vehicle.get('co2_g_km')} g/km ({vehicle.get('co2_norme') or '-'})" if vehicle.get("co2_g_km") else "-"),
        ("Conso officielle", f"{vehicle.get('conso_officielle_l_100km')} L/100 km ({vehicle.get('conso_officielle_norme') or '-'})"
         if vehicle.get("conso_officielle_l_100km") else "-"),
        ("Conso réelle", f"{reelle} L/100 km (source : {sources.get(vehicle.get('conso_reelle_source'), '-')})"
         if reelle else "Données insuffisantes"),
        ("Capacité réservoir", f"{vehicle.get('capacite_reservoir_l')} L" if vehicle.get("capacite_reservoir_l") else "-"),
    ])
    pdf.ln(4)

    ml = m.get("leasing") or {}
    _section(pdf, "Leasing")
    _kv_table(pdf, [
        ("Société / Contrat", f"{leasing.get('societe') or '-'} / {leasing.get('numero_contrat') or '-'}"),
        ("Période", f"{_date_fr(leasing.get('date_debut'))} -> {_due(leasing.get('date_fin'), ml.get('days_remaining'))}"),
        ("Mensualité", _chf(leasing.get("mensualite_chf") or leasing.get("cout_mensuel"))),
        ("Mois restants / Coût restant", f"{ml.get('months_remaining') if ml.get('months_remaining') is not None else '-'} / {_chf(ml.get('cost_remaining'))}"),
        ("Km contractuel / annuel", f"{leasing.get('km_contractuel') or '-'} / {leasing.get('km_annuel') or '-'}"),
    ])
    pdf.ln(4)

    _section(pdf, "Assurance")
    _kv_table(pdf, [
        ("Compagnie / Police", f"{assurance.get('compagnie') or '-'} / {assurance.get('numero_police') or '-'}"),
        ("Couverture", assurance.get("type_couverture")),
        ("Échéance", _due(assurance.get("date_echeance"), (m.get("assurance") or {}).get("days_remaining"))),
        ("Prime annuelle / Franchise", f"{_chf(assurance.get('prime_annuelle'))} / {_chf(assurance.get('franchise'))}"),
    ])
    pdf.ln(4)

    _section(pdf, "Contrôle technique")
    _kv_table(pdf, [
        ("Dernier contrôle", _date_fr(controle.get("date_dernier")) if controle.get("date_dernier") else "-"),
        ("Prochain contrôle", _due(controle.get("date_prochain"), (m.get("controle") or {}).get("days_remaining"))),
        ("Centre / Résultat", f"{controle.get('centre') or '-'} / {controle.get('resultat') or '-'}"),
    ])
    pdf.ln(4)

    _section(pdf, f"Documents ({len(documents)})")
    if documents:
        pdf.set_font("helvetica", "", 8)
        with pdf.table(col_widths=(38, 20, 22, 12, 8), text_align=("LEFT", "LEFT", "LEFT", "LEFT", "RIGHT"),
                       line_height=5.5, padding=1) as table:
            head = table.row()
            for h in ("Nom du fichier", "Dossier", "Type", "Ajouté le", "Taille"):
                head.cell(_tx(h))
            for d in documents:
                row = table.row()
                row.cell(_tx(d.get("original_filename") or "-"))
                row.cell(_tx(d.get("folder") or "-"))
                row.cell(_tx(labels.get(d.get("document_type"), d.get("document_type") or "-")))
                row.cell(_tx(_date_fr(d.get("created_at"))))
                row.cell(_tx(_size_h(d.get("size"))))
    else:
        pdf.set_font("helvetica", "I", 9)
        pdf.cell(0, 6, _tx("Aucun document."), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    _section(pdf, f"Historique ({len(history)} dernier(s) événement(s))")
    if history:
        pdf.set_font("helvetica", "", 8)
        with pdf.table(col_widths=(18, 14, 68), text_align=("LEFT", "LEFT", "LEFT"),
                       line_height=5.5, padding=1) as table:
            head = table.row()
            for h in ("Date", "Action", "Détail"):
                head.cell(_tx(h))
            for e in history:
                row = table.row()
                row.cell(_tx(_dt_fr(e.get("created_at"))))
                row.cell(_tx(e.get("action") or "-"))
                row.cell(_tx(e.get("detail") or "-"))
    else:
        pdf.set_font("helvetica", "I", 9)
        pdf.cell(0, 6, _tx("Aucun événement enregistré."), new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# Phase 4C — Lot D : exports Amendes (CSV / XLSX / PDF). Lignes = périmètre filtré de /api/fines ;
# `notes_internes` jamais exporté ; `paid_on` = date métier (jamais `paid_at`) ; `payment_ref` seulement si présente.
# ---------------------------------------------------------------------------
FINE_EXPORT_COLUMNS = [
    ("Référence", lambda d: d.get("numero") or ""),
    ("Dossier interne", lambda d: d.get("dossier_interne") or ""),
    ("Véhicule", lambda d: d.get("plaque") or ""),
    ("Marque / modèle", lambda d: d.get("vehicule_label") or ""),
    ("Conducteur", lambda d: d.get("driver_nom") or ""),
    ("Statut", lambda d: d.get("fine_status_label") or d.get("fine_status") or ""),
    ("Type d'infraction", lambda d: d.get("type_infraction_label") or d.get("type_infraction") or ""),
    ("Autorité", lambda d: d.get("fournisseur") or ""),
    ("Lieu", lambda d: d.get("lieu_label") or ""),
    ("Date infraction", lambda d: _date_fr(d.get("date_debut")) if d.get("date_debut") else ""),
    ("Échéance", lambda d: _date_fr(d.get("date_expiration")) if d.get("date_expiration") else ""),
    ("Montant", lambda d: _n(d.get("montant")) if d.get("montant") is not None else ""),
    ("Devise", lambda d: d.get("devise") or "CHF"),
    ("Montant CHF", lambda d: "conversion en attente" if d.get("pending_fx") else (_n(d.get("montant_chf_effectif")) if d.get("montant_chf_effectif") is not None else "")),
    ("Frais admin", lambda d: _n(d.get("frais_admin")) if d.get("frais_admin") is not None else ""),
    ("Priorité", lambda d: d.get("priorite") or ""),
    ("Payée le", lambda d: _date_fr(d.get("paid_on")) if d.get("paid_on") else ""),
    ("Réf. paiement", lambda d: d.get("payment_ref") or ""),
    ("Sans justificatif", lambda d: "oui" if d.get("justificatif_absent") else ""),
    ("Pièces liées", lambda d: d.get("attachments_count") or 0),
]


def build_fines_csv(rows: list) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow([c for c, _ in FINE_EXPORT_COLUMNS])
    for d in rows:
        w.writerow([f(d) for _, f in FINE_EXPORT_COLUMNS])
    return buf.getvalue()


def build_fines_xlsx(rows: list, totals: dict, meta: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Amendes"
    ws.append([c for c, _ in FINE_EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
    for d in rows:
        ws.append([f(d) for _, f in FINE_EXPORT_COLUMNS])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"
    s = wb.create_sheet("Synthèse")
    s.append(["Généré le", _dt_fr(meta.get("generated_at"))])
    s.append(["Lignes", totals.get("count", len(rows))])
    s.append(["Total compté CHF (hors annulées, hors conversion en attente)", totals.get("total_chf", 0)])
    s.append(["Ouvert CHF (échéances actives)", totals.get("ouvert_chf", 0)])
    s.append(["Payé CHF", totals.get("paye_chf", 0)])
    s.append(["Annulé CHF (montant conservé, exclu du total)", totals.get("annule_chf", 0)])
    s.append(["Conversion en attente (lignes)", totals.get("pending_fx_count", 0)])
    s.append(["Filtres", ", ".join(f"{k}={v}" for k, v in (meta.get("filters") or {}).items()) or "aucun"])
    for st, n in (totals.get("by_status") or {}).items():
        s.append([f"Statut {st}", n])
    s.column_dimensions["A"].width = 62
    s.column_dimensions["B"].width = 24
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class _FinesReport(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("helvetica", "I", 8)
        self.set_text_color(130)
        self.cell(0, 8, _tx(f"LogiTrak Documents — Amendes — page {self.page_no()}"), align="C")
        self.set_text_color(0)


def build_fines_pdf(rows: list, totals: dict, meta: dict) -> bytes:
    pdf = _FinesReport(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_font("helvetica", "B", 16)
    pdf.cell(0, 9, _tx("Amendes — export"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 9)
    pdf.set_text_color(100)
    filters = ", ".join(f"{k}={v}" for k, v in (meta.get("filters") or {}).items()) or "aucun filtre"
    pdf.cell(0, 5, _tx(f"Généré le {_dt_fr(meta.get('generated_at'))} UTC · {len(rows)} amende(s) · Filtres : {filters}"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)
    _section(pdf, "Synthèse")
    pdf.set_font("helvetica", "", 10)
    pdf.cell(0, 6, _tx(f"Total compté : {totals.get('total_chf', 0):.2f} CHF   ·   Ouvert : {totals.get('ouvert_chf', 0):.2f} CHF   ·   "
                       f"Payé : {totals.get('paye_chf', 0):.2f} CHF   ·   Annulé (exclu) : {totals.get('annule_chf', 0):.2f} CHF   ·   "
                       f"Conversion en attente : {totals.get('pending_fx_count', 0)}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    _section(pdf, "Détail")
    pdf.set_font("helvetica", "", 7.5)
    cols = [("Référence", 30), ("Véhicule", 22), ("Conducteur", 30), ("Statut", 26), ("Type", 24), ("Autorité", 36),
            ("Infraction", 20), ("Échéance", 20), ("Montant", 22), ("Payée le", 20), ("Réf. paiement", 27)]
    getters = {c: f for c, f in FINE_EXPORT_COLUMNS}
    with pdf.table(col_widths=tuple(w for _, w in cols), line_height=5.5, padding=1) as table:
        head = table.row()
        for c, _ in cols:
            head.cell(_tx(c))
        for d in rows:
            row = table.row()
            row.cell(_tx(getters["Référence"](d) or d.get("dossier_interne") or "-"))
            row.cell(_tx(getters["Véhicule"](d) or "-"))
            row.cell(_tx(getters["Conducteur"](d) or "-"))
            row.cell(_tx(getters["Statut"](d)))
            row.cell(_tx(getters["Type d'infraction"](d) or "-"))
            row.cell(_tx(getters["Autorité"](d) or "-"))
            row.cell(_tx(getters["Date infraction"](d) or "-"))
            row.cell(_tx(getters["Échéance"](d) or "-"))
            m = d.get("montant")
            row.cell(_tx(f"{m:.2f} {d.get('devise') or 'CHF'}" if m is not None else "-"))
            row.cell(_tx(getters["Payée le"](d) or "-"))
            row.cell(_tx(getters["Réf. paiement"](d) or "-"))
    return bytes(pdf.output())
