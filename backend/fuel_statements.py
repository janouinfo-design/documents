"""Phase 4C — Lot G (6b) : rapprochement achats ↔ consommation, décomptes périodiques (statements) clôturables/verrouillants, exports.
Fonctions pures (testables sans DB). Décisions figées : CAN mesuré = consommation réelle prioritaire ; tickets/achats ≠ consommation réelle
(indicatif seulement) ; ASTRA = référence comparative, n'écrase jamais CAN ; seuils tenant null par défaut → statut INDICATIF ;
statuts OK | A_CONTROLER | INDICATIF | IMPOSSIBLE ; mois civil Europe/Zurich (D5) ; blockers de ligne = pending_fx | matched_review |
unmatched | open_anomaly | forced_duplicate ; verrou uniquement à la clôture (normale ou exception) ; aucune réouverture (correctif seulement).
Règle de seuils (documentée, jamais implicite) : un seul seuil configuré → il s'applique seul ; deux seuils → À CONTRÔLER seulement si
l'écart dépasse LES DEUX tolérances (% ET litres)."""
import csv
import hashlib
import io
import json
import re
from datetime import date, datetime
from typing import Optional
from zoneinfo import ZoneInfo

from pydantic import BaseModel

TZ = ZoneInfo("Europe/Zurich")
PERIOD_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
RECO_STATUSES = ("OK", "A_CONTROLER", "INDICATIF", "IMPOSSIBLE")
RECO_LABELS = {"OK": "Cohérent", "A_CONTROLER": "Écart à contrôler", "INDICATIF": "Indicatif", "IMPOSSIBLE": "Impossible"}
STATEMENT_STATUSES = ("brouillon", "cloture")
STATEMENT_LABELS = {"brouillon": "Brouillon", "cloture": "Clôturé"}
STATEMENT_TYPES = ("regulier", "correctif")
BLOCKER_TYPES = ("pending_fx", "matched_review", "unmatched", "open_anomaly", "forced_duplicate")
BLOCKER_LABELS = {"pending_fx": "Conversion CHF en attente", "matched_review": "Véhicule à vérifier", "unmatched": "Véhicule non rattaché",
                  "open_anomaly": "Anomalie ouverte", "forced_duplicate": "Doublon forcé"}
THRESHOLD_RULE = {"code": "single_or_both", "label": "Un seul seuil configuré : il s'applique seul · deux seuils : écart À CONTRÔLER seulement "
                                                     "s'il dépasse les deux tolérances (% ET litres) · aucun seuil : INDICATIF"}
REASON_MIN_LEN = 3
LOCK_CODE = "STATEMENT_LOCKED"
# Champs de DocumentUpdate utilisés (directement ou via la fuel_transaction) par le snapshot d'un décompte
DOC_PROTECTED_FIELDS = ("montant", "devise", "montant_chf", "date_debut", "date_expiration", "fournisseur", "numero", "driver_id",
                        "kilometrage_releve", "montant_ht", "tva_chf", "business_category", "frequence")


# --- période ---------------------------------------------------------------------------------------------------------
def valid_period(p: Optional[str]) -> bool:
    return bool(p) and bool(PERIOD_RE.match(p))


def period_bounds(period_month: str) -> tuple:
    y, m = int(period_month[:4]), int(period_month[5:7])
    first = date(y, m, 1)
    last = date(y + (m == 12), (m % 12) + 1, 1).toordinal() - 1
    return first.isoformat(), date.fromordinal(last).isoformat()


def current_period() -> str:
    return datetime.now(TZ).strftime("%Y-%m")


def shift_period(period_month: str, delta: int) -> str:
    y, m = int(period_month[:4]), int(period_month[5:7]) + delta
    y, m = y + (m - 1) // 12, (m - 1) % 12 + 1
    return f"{y:04d}-{m:02d}"


def tx_day(tx: dict) -> Optional[str]:
    """Jour métier de la transaction (Europe/Zurich, D5) : `date` (déjà locale) sinon `date_heure` converti."""
    if tx.get("date"):
        return str(tx["date"])[:10]
    dh = tx.get("date_heure")
    if not dh:
        return None
    try:
        d = datetime.fromisoformat(str(dh).replace("Z", "+00:00"))
        return (d.astimezone(TZ) if d.tzinfo else d).date().isoformat()
    except ValueError:
        return str(dh)[:10]


def in_period(tx: dict, date_from: str, date_to: str) -> bool:
    d = tx_day(tx)
    return bool(d) and date_from <= d <= date_to


def in_scope(tx: dict, scope: dict) -> bool:
    if (scope or {}).get("type") == "fournisseur":
        return (tx.get("fournisseur") or "").strip().lower() == (scope.get("fournisseur") or "").strip().lower()
    return True


# --- paramètres -----------------------------------------------------------------------------------------------------
def reconciliation_settings(doc: Optional[dict]) -> dict:
    r = ((doc or {}).get("fuel") or {}).get("reconciliation") or {}
    return {"threshold_pct": r.get("threshold_pct"), "threshold_l": r.get("threshold_l")}


class ReconciliationSettingsPayload(BaseModel):
    threshold_pct: Optional[float] = None
    threshold_l: Optional[float] = None


def reconciliation_settings_errors(p: ReconciliationSettingsPayload) -> list:
    return [f"{k} : nombre > 0 ou null" for k in ("threshold_pct", "threshold_l") if getattr(p, k) is not None and getattr(p, k) <= 0]


# --- rapprochement ---------------------------------------------------------------------------------------------------
def tx_chf(tx: dict) -> Optional[float]:
    """D7 : CHF → montant ; devise ≠ CHF + montant_chf → montant_chf ; sinon None (conversion en attente, exclu)."""
    if tx.get("montant") is None:
        return None
    if (tx.get("devise") or "CHF") == "CHF":
        return float(tx["montant"])
    return float(tx["montant_chf"]) if tx.get("montant_chf") is not None else None


def purchases(txs: list) -> dict:
    out = {"litres": 0.0, "kwh": 0.0, "chf": 0.0, "n_tx": 0, "pending_fx": 0, "by_fournisseur": {}}
    for t in txs:
        out["n_tx"] += 1
        out["litres"] += float(t.get("litres") or 0) if t.get("energie") != "electrique" else 0.0
        out["kwh"] += float(t.get("energie_kwh") or 0)
        chf = tx_chf(t)
        if chf is None:
            out["pending_fx"] += 1
        else:
            out["chf"] += chf
        f = t.get("fournisseur") or "—"
        out["by_fournisseur"][f] = out["by_fournisseur"].get(f, 0) + 1
    for k in ("litres", "kwh", "chf"):
        out[k] = round(out[k], 2)
    return out


def can_consumption(snaps: list, date_from: str, date_to: str) -> Optional[dict]:
    """Consommation MESURÉE CAN sur la période : Δ litres cumulés entre le dernier relevé avant/au début de période et le dernier relevé
    de la période (fuel_snapshots{day, litres_cumules, km}). None si données insuffisantes — jamais 0 inventé."""
    pts = sorted([s for s in snaps if s.get("day") and s.get("litres_cumules") is not None and s.get("km") is not None], key=lambda s: s["day"])
    inside = [s for s in pts if date_from <= s["day"] <= date_to]
    if not inside:
        return None
    end = inside[-1]
    before = [s for s in pts if s["day"] < date_from]
    start = before[-1] if before else inside[0]
    if start is end:
        return None
    dl, dkm = float(end["litres_cumules"]) - float(start["litres_cumules"]), float(end["km"]) - float(start["km"])
    if dkm <= 0 or dl < 0:
        return None
    return {"source": "can", "litres": round(dl, 2), "km": round(dkm), "l_100km": round(dl / dkm * 100, 1) if dkm >= 100 else None,
            "from_day": start["day"], "to_day": end["day"], "snapshots": len(inside) + (1 if before else 0)}


def tickets_estimate(txs: list) -> Optional[dict]:
    """ESTIMATION par tickets (Σ litres des pleins suivants / Δ km ≥ 100) — achats ≠ consommation réelle : indicatif uniquement."""
    pts = sorted([t for t in txs if t.get("energie") != "electrique" and t.get("kilometrage") and t.get("litres")],
                 key=lambda t: (float(t["kilometrage"]), t.get("date_heure") or ""))
    if len(pts) < 2:
        return None
    dkm = float(pts[-1]["kilometrage"]) - float(pts[0]["kilometrage"])
    if dkm < 100:
        return None
    litres = sum(float(t["litres"]) for t in pts[1:])
    return {"source": "tickets", "litres": round(litres, 2), "km": round(dkm), "l_100km": round(litres / dkm * 100, 1), "n": len(pts)}


def threshold_verdict(ecart_l: Optional[float], ecart_pct: Optional[float], th_pct: Optional[float], th_l: Optional[float]) -> Optional[str]:
    """None si aucun seuil ; sinon OK / A_CONTROLER selon THRESHOLD_RULE (un seuil : seul ; deux seuils : les deux dépassés)."""
    checks = []
    if th_pct is not None:
        checks.append(ecart_pct is not None and abs(ecart_pct) > th_pct)
    if th_l is not None:
        checks.append(ecart_l is not None and abs(ecart_l) > th_l)
    if not checks:
        return None
    return "A_CONTROLER" if all(checks) else "OK"


def reconcile(vehicle: dict, txs: list, snaps: list, settings: dict, period_month: str, open_anomalies: Optional[dict] = None,
              justification: Optional[dict] = None) -> dict:
    """Rapprochement achats ↔ consommation d'un véhicule sur un mois. CAN mesuré = consommation réelle ; tickets = estimation indicative ;
    ASTRA = référence comparative. Statut : IMPOSSIBLE (ni consommation ni achat) · INDICATIF (tickets seuls, ou CAN sans seuil) ·
    OK / A_CONTROLER (CAN + seuils, THRESHOLD_RULE). La justification explique l'écart, ne le corrige jamais."""
    date_from, date_to = period_bounds(period_month)
    ach = purchases(txs)
    can = can_consumption(snaps, date_from, date_to)
    est = tickets_estimate(txs)
    conso = can or {"source": "unavailable", "litres": None, "km": None, "l_100km": None}
    ecart_l = round(ach["litres"] - conso["litres"], 2) if conso["litres"] is not None else None
    ecart_pct = round(ecart_l / conso["litres"] * 100, 1) if ecart_l is not None and conso["litres"] else None
    th_pct, th_l = settings.get("threshold_pct"), settings.get("threshold_l")
    verdict = threshold_verdict(ecart_l, ecart_pct, th_pct, th_l)
    if ach["n_tx"] == 0 and can is None:
        status, reason = "IMPOSSIBLE", "aucun achat ni mesure CAN sur la période"
    elif can is None:
        status, reason = "INDICATIF", ("achats sans mesure CAN indépendante — estimation tickets indicative (achats ≠ consommation réelle)" if est
                                       else "achats sans mesure CAN : aucune consommation réelle disponible")
    elif ach["n_tx"] == 0:
        status, reason = "INDICATIF", "mesure CAN sans aucun achat enregistré sur la période"
    elif verdict is None:
        status, reason = "INDICATIF", "aucun seuil configuré (threshold_pct / threshold_l null) : écart brut affiché, aucun verdict"
    elif verdict == "A_CONTROLER":
        status, reason = "A_CONTROLER", f"écart {ecart_l:+.2f} L ({ecart_pct:+.1f} %) au-delà des seuils configurés ({_th_label(th_pct, th_l)})"
    else:
        status, reason = "OK", f"écart {ecart_l:+.2f} L ({ecart_pct:+.1f} %) dans les seuils configurés ({_th_label(th_pct, th_l)})"
    ref = vehicle.get("conso_officielle_l_100km")
    ref = float(ref) if ref not in (None, "", 0) else None
    real = conso.get("l_100km")
    astra = {"conso_officielle_l_100km": ref, "norme": vehicle.get("conso_officielle_norme") or None, "conso_reelle_l_100km": real,
             "ecart_l_100km": round(real - ref, 1) if ref is not None and real is not None else None,
             "ecart_pct": round((real - ref) / ref * 100, 1) if ref and real is not None else None, "role": "reference_comparative"}
    blockers = {}
    for t in txs:
        for b in line_blockers(t, (open_anomalies or {}).get(t["id"], 0)):
            blockers[b] = blockers.get(b, 0) + 1
    return {"vehicle_id": vehicle["id"], "plaque": vehicle.get("plaque"), "vehicule_label": " ".join(x for x in (vehicle.get("marque"), vehicle.get("modele")) if x) or None,
            "period_month": period_month, "period_from": date_from, "period_to": date_to, "achats": ach, "consommation": conso, "estimation_tickets": est,
            "source_consumption": conso["source"], "ecart_l": ecart_l, "ecart_pct": ecart_pct, "status": status, "status_label": RECO_LABELS[status],
            "status_reason": reason, "thresholds": {"threshold_pct": th_pct, "threshold_l": th_l, "configured": verdict is not None, "rule": THRESHOLD_RULE},
            "astra": astra, "blockers": {"count": sum(blockers.values()), "by_type": blockers}, "justification": justification or None,
            "transaction_ids": [t["id"] for t in txs],
            "breakdown": [{"step": "achats", "detail": f"{ach['n_tx']} transaction(s) · {ach['litres']} L · {ach['kwh']} kWh · {ach['chf']} CHF comptés ({ach['pending_fx']} en attente FX)"},
                          {"step": "consommation", "detail": f"CAN mesuré {can['litres']} L sur {can['km']} km ({can['from_day']} → {can['to_day']})" if can else "aucune mesure CAN disponible"},
                          {"step": "estimation_tickets", "detail": f"tickets {est['litres']} L / {est['km']} km (estimation indicative, pas une consommation réelle)" if est else "non calculable"},
                          {"step": "astra", "detail": f"référence {ref} L/100 km ({astra['norme'] or 'norme ?'})" if ref is not None else "aucune référence ASTRA"},
                          {"step": "ecart", "detail": f"achats − consommation CAN = {ecart_l} L ({ecart_pct} %)" if ecart_l is not None else "non calculable"},
                          {"step": "statut", "detail": reason}]}


def _th_label(th_pct, th_l) -> str:
    return " et ".join(x for x in ((f"{th_pct} %" if th_pct is not None else None), (f"{th_l} L" if th_l is not None else None)) if x)


class JustifyPayload(BaseModel):
    reason: str


# --- décomptes -------------------------------------------------------------------------------------------------------
class StatementCreate(BaseModel):
    period_month: Optional[str] = None
    scope_type: str = "tenant"
    fournisseur: Optional[str] = None
    type: str = "regulier"
    reason: Optional[str] = None
    parent_statement_id: Optional[str] = None


class DeclaredPayload(BaseModel):
    montant: Optional[float] = None
    devise: Optional[str] = None
    volume_l: Optional[float] = None
    kwh: Optional[float] = None
    nb_lignes: Optional[int] = None


class ReasonPayload(BaseModel):
    reason: str


class CloseExceptionPayload(BaseModel):
    reason: str
    confirm: bool = False


def statement_create_errors(p: StatementCreate) -> list:
    e = []
    if p.type not in STATEMENT_TYPES:
        e.append("type ∈ {regulier, correctif}")
    if p.type == "correctif":
        if not (p.parent_statement_id or "").strip():
            e.append("parent_statement_id obligatoire pour un décompte correctif")
        if len((p.reason or "").strip()) < REASON_MIN_LEN:
            e.append("motif obligatoire pour un décompte correctif (≥ 3 caractères)")
        return e
    if not valid_period(p.period_month):
        e.append("period_month au format YYYY-MM")
    if p.scope_type not in ("tenant", "fournisseur"):
        e.append("scope_type ∈ {tenant, fournisseur}")
    if p.scope_type == "fournisseur" and not (p.fournisseur or "").strip():
        e.append("fournisseur obligatoire pour scope fournisseur")
    return e


def declared_errors(p: DeclaredPayload) -> list:
    e = []
    for k in ("montant", "volume_l", "kwh"):
        if getattr(p, k) is not None and getattr(p, k) < 0:
            e.append(f"{k} ≥ 0")
    if p.nb_lignes is not None and p.nb_lignes < 0:
        e.append("nb_lignes ≥ 0")
    if p.devise is not None and not re.fullmatch(r"[A-Za-z]{3}", p.devise.strip()):
        e.append("devise : code ISO 3 lettres")
    return e


def declared_from(p: DeclaredPayload) -> dict:
    return {"montant": p.montant, "devise": p.devise.strip().upper() if p.devise else None, "volume_l": p.volume_l, "kwh": p.kwh, "nb_lignes": p.nb_lignes}


def scope_of(p: StatementCreate) -> dict:
    return {"type": "fournisseur", "fournisseur": p.fournisseur.strip()} if p.scope_type == "fournisseur" else {"type": "tenant", "fournisseur": None}


def scope_label(scope: dict) -> str:
    return f"Fournisseur {scope.get('fournisseur')}" if (scope or {}).get("type") == "fournisseur" else "Tout le tenant"


def line_blockers(tx: dict, open_anomalies: int) -> list:
    b = []
    if tx_chf(tx) is None and tx.get("montant") is not None:
        b.append("pending_fx")
    if tx.get("match_status") in ("matched_review", "unmatched"):
        b.append(tx["match_status"])
    if open_anomalies:
        b.append("open_anomaly")
    if tx.get("forced_duplicate_of"):
        b.append("forced_duplicate")
    return b


def build_line(tx: dict, open_anomalies: int, plaque: Optional[str]) -> dict:
    line = {"transaction_id": tx["id"], "document_id": tx.get("source_document_id"), "vehicle_id": tx.get("vehicle_id"), "plaque": plaque,
            "date": tx_day(tx), "heure": tx.get("heure"), "fournisseur": tx.get("fournisseur"), "station": tx.get("station"), "carte_last4": tx.get("carte_last4"),
            "card_id": tx.get("card_id"), "driver_id": tx.get("driver_id"), "montant": tx.get("montant"), "devise": tx.get("devise") or "CHF", "montant_chf": tx_chf(tx),
            "litres": tx.get("litres") if tx.get("energie") != "electrique" else None, "kwh": tx.get("energie_kwh"), "kilometrage": tx.get("kilometrage"),
            "energie": tx.get("energie"), "type_carburant": tx.get("type_carburant"), "match_status": tx.get("match_status"),
            "forced_duplicate_of": tx.get("forced_duplicate_of"), "open_anomalies": open_anomalies, "blockers": line_blockers(tx, open_anomalies),
            "external_transaction_id": tx.get("external_transaction_id"), "created_from": tx.get("created_from"), "tx_created_at": tx.get("created_at")}
    line["fingerprint"] = line_fingerprint(line)
    return line


_FP_KEYS = ("transaction_id", "document_id", "vehicle_id", "date", "heure", "fournisseur", "station", "carte_last4", "card_id", "driver_id", "montant", "devise",
            "montant_chf", "litres", "kwh", "kilometrage", "energie", "type_carburant", "match_status", "forced_duplicate_of", "open_anomalies", "blockers")


def line_fingerprint(line: dict) -> str:
    """Empreinte des données du snapshot utilisées par le décompte : toute divergence avec l'état courant = snapshot obsolète (recalcul explicite requis)."""
    return hashlib.sha256(json.dumps({k: line.get(k) for k in _FP_KEYS}, sort_keys=True, default=str).encode()).hexdigest()


def totals_of(lines: list) -> dict:
    t = {"n_lignes": len(lines), "montant_chf": 0.0, "litres": 0.0, "kwh": 0.0, "pending_fx": 0, "blocker_count": 0, "blocked_line_count": 0,
         "blockers_by_type": {}, "n_vehicules": len({ln.get("vehicle_id") for ln in lines if ln.get("vehicle_id")})}
    for ln in lines:
        if ln["montant_chf"] is None:
            t["pending_fx"] += 1
        else:
            t["montant_chf"] += ln["montant_chf"]
        t["litres"] += float(ln.get("litres") or 0)
        t["kwh"] += float(ln.get("kwh") or 0)
        if ln["blockers"]:
            t["blocked_line_count"] += 1
            t["blocker_count"] += len(ln["blockers"])
            for b in ln["blockers"]:
                t["blockers_by_type"][b] = t["blockers_by_type"].get(b, 0) + 1
    for k in ("montant_chf", "litres", "kwh"):
        t[k] = round(t[k], 2)
    return t


def declared_deltas(totals: dict, declared: Optional[dict]) -> dict:
    """Écarts Documents − relevé fournisseur déclaré ; N/A (None) si valeur déclarée absente ; devise ≠ CHF → non comparable (D7)."""
    d = declared or {}
    out = {"delta_montant": None, "delta_montant_pct": None, "montant_comparable": None, "delta_volume_l": None, "delta_kwh": None, "delta_nb_lignes": None}
    if d.get("montant") is not None:
        if (d.get("devise") or "CHF").upper() == "CHF":
            out["montant_comparable"] = True
            out["delta_montant"] = round(totals["montant_chf"] - float(d["montant"]), 2)
            out["delta_montant_pct"] = round(out["delta_montant"] / float(d["montant"]) * 100, 1) if float(d["montant"]) else None
        else:
            out["montant_comparable"] = False
    if d.get("volume_l") is not None:
        out["delta_volume_l"] = round(totals["litres"] - float(d["volume_l"]), 2)
    if d.get("kwh") is not None:
        out["delta_kwh"] = round(totals["kwh"] - float(d["kwh"]), 2)
    if d.get("nb_lignes") is not None:
        out["delta_nb_lignes"] = totals["n_lignes"] - int(d["nb_lignes"])
    return out


def statement_number(period_month: str, seq: int, stype: str) -> str:
    return f"{'COR' if stype == 'correctif' else 'DEC'}-{period_month}-{seq:03d}"


def blockers_snapshot(lines: list) -> dict:
    t = totals_of(lines)
    return {"blocker_count": t["blocker_count"], "blocked_line_count": t["blocked_line_count"], "blockers_by_type": t["blockers_by_type"],
            "lines": [{"transaction_id": ln["transaction_id"], "plaque": ln.get("plaque"), "date": ln.get("date"), "montant": ln.get("montant"),
                       "devise": ln.get("devise"), "blockers": ln["blockers"]} for ln in lines if ln["blockers"]]}


def integrity_errors(lines: list, current_lines: dict, statement_id: str) -> list:
    """Contrôle d'intégrité avant clôture : chaque ligne du snapshot doit correspondre à une transaction existante, non supprimée,
    non verrouillée par un autre décompte, et dont les données courantes ont la même empreinte (sinon snapshot obsolète → recalcul explicite)."""
    e = []
    for ln in lines:
        cur = current_lines.get(ln["transaction_id"])
        if not cur:
            e.append({"code": "TX_MISSING", "transaction_id": ln["transaction_id"], "detail": "transaction supprimée ou introuvable"})
        elif cur.get("locked") and cur.get("statement_id") != statement_id:
            e.append({"code": "TX_LOCKED_ELSEWHERE", "transaction_id": ln["transaction_id"], "detail": f"déjà verrouillée par le décompte {cur.get('statement_id')}"})
        elif cur["line"]["fingerprint"] != ln.get("fingerprint"):
            e.append({"code": "SNAPSHOT_STALE", "transaction_id": ln["transaction_id"], "detail": "données ou blockers modifiés depuis le snapshot — recalcul explicite requis"})
    return e


# --- exports génériques (CSV / XLSX / PDF) ----------------------------------------------------------------------------
def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "oui" if v else "non"
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}={val}" for k, val in v.items())
    return v


def export_csv(columns: list, rows: list, summary: list = None) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    for k, v in summary or []:
        w.writerow([f"# {k}", _cell(v)])
    w.writerow([c for c, _ in columns])
    for r in rows:
        w.writerow([_cell(r.get(k)) for _, k in columns])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def export_xlsx(columns: list, rows: list, title: str, summary: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\\/*?:\[\]]", "-", title)[:31]
    ws.append([c for c, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
    for r in rows:
        ws.append([_cell(r.get(k)) for _, k in columns])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(10, min(42, max(len(str(c.value or "")) for c in col) + 2))
    ws.freeze_panes = "A2"
    s = wb.create_sheet("Synthèse")
    for k, v in summary:
        s.append([k, _cell(v)])
    s.column_dimensions["A"].width = 56
    s.column_dimensions["B"].width = 30
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_pdf(columns: list, rows: list, title: str, summary: list) -> bytes:
    from fpdf import FPDF
    tx = lambda s: str(s).encode("latin-1", "replace").decode("latin-1")  # noqa: E731
    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    pdf.set_font("helvetica", "B", 15)
    pdf.cell(0, 9, tx(title), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("helvetica", "", 9)
    for k, v in summary:
        pdf.cell(0, 5, tx(f"{k} : {_cell(v)}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    width = (pdf.w - 2 * pdf.l_margin) / max(1, len(columns))
    pdf.set_font("helvetica", "B", 7)
    for c, _ in columns:
        pdf.cell(width, 6, tx(str(c)[:22]), border=1)
    pdf.ln()
    pdf.set_font("helvetica", "", 7)
    for r in rows:
        for _, k in columns:
            pdf.cell(width, 5.5, tx(str(_cell(r.get(k)))[:26]), border=1)
        pdf.ln()
    if not rows:
        pdf.cell(0, 6, tx("Aucune ligne."), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


EXPORT_MIME = {"csv": "text/csv; charset=utf-8", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf"}


def export(fmt: str, columns: list, rows: list, title: str, summary: list) -> tuple:
    """→ (bytes, mime, sha256) — le hash est calculé sur les octets finaux réellement renvoyés."""
    if fmt == "csv":
        data = export_csv(columns, rows, summary)
    elif fmt == "xlsx":
        data = export_xlsx(columns, rows, title, summary)
    elif fmt == "pdf":
        data = export_pdf(columns, rows, title, summary)
    else:
        raise ValueError("fmt ∈ {csv, xlsx, pdf}")
    return data, EXPORT_MIME[fmt], hashlib.sha256(data).hexdigest()


STATEMENT_COLUMNS = [("Date", "date"), ("Heure", "heure"), ("Véhicule", "plaque"), ("Fournisseur", "fournisseur"), ("Station", "station"), ("Carte", "carte_last4"),
                     ("Montant", "montant"), ("Devise", "devise"), ("Montant CHF", "montant_chf"), ("Litres", "litres"), ("kWh", "kwh"),
                     ("Rattachement", "match_status"), ("Anomalies ouvertes", "open_anomalies"), ("Blockers", "blockers"), ("Verrouillée", "locked"),
                     ("Transaction", "transaction_id"), ("Document", "document_id")]
RECO_COLUMNS = [("Période", "period_month"), ("Véhicule", "plaque"), ("Modèle", "vehicule_label"), ("Achats L", "achats_litres"), ("Achats kWh", "achats_kwh"), ("Achats CHF", "achats_chf"),
                ("Transactions", "achats_n_tx"), ("Conso source", "conso_source"), ("Conso CAN L", "conso_litres"), ("Km", "conso_km"), ("Conso réelle L/100", "conso_l_100km"),
                ("ASTRA L/100", "astra_ref"), ("Écart ASTRA L/100", "astra_ecart"), ("Écart L", "ecart_l"), ("Écart %", "ecart_pct"), ("Statut", "status"), ("Motif statut", "status_reason"),
                ("Blockers", "blockers"), ("Justification", "justification")]
ENERGY_COLUMNS = [("Date", "date"), ("Heure", "heure"), ("Véhicule", "plaque"), ("Conducteur", "driver_nom"), ("Fournisseur", "fournisseur"), ("Station", "station"), ("Carte", "carte_last4"),
                  ("Énergie", "energie"), ("Produit", "type_carburant"), ("Litres", "litres"), ("kWh", "energie_kwh"), ("Prix unitaire", "prix_unitaire"), ("Montant", "montant"), ("Devise", "devise"),
                  ("Montant CHF", "montant_chf"), ("Km", "kilometrage"), ("Rattachement", "match_status"), ("Anomalies ouvertes", "anomalies_open"), ("Verrouillée", "locked"),
                  ("Décompte", "statement_number"), ("Source", "created_from"), ("Transaction", "id")]


def reco_row(r: dict) -> dict:
    a, c, s = r["achats"], r["consommation"], r.get("astra") or {}
    j = r.get("justification") or {}
    return {**{k: r.get(k) for k in ("period_month", "plaque", "vehicule_label", "ecart_l", "ecart_pct", "status", "status_reason")},
            "achats_litres": a["litres"], "achats_kwh": a["kwh"], "achats_chf": a["chf"], "achats_n_tx": a["n_tx"], "conso_source": c["source"], "conso_litres": c.get("litres"),
            "conso_km": c.get("km"), "conso_l_100km": c.get("l_100km"), "astra_ref": s.get("conso_officielle_l_100km"), "astra_ecart": s.get("ecart_l_100km"),
            "blockers": (r.get("blockers") or {}).get("by_type") or None,
            "justification": f"{j.get('reason')} ({j.get('by')}, {str(j.get('at') or '')[:10]})" if j.get("reason") else None}
