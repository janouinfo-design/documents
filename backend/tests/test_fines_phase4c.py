"""Phase 4C — Lot D : amendes — 10 statuts (parité Journal), paiement métier (paid_on ≠ paid_at, payment_ref),
dé-paiement motivé + audit `fine_payment_reverted`, D9 (annulee hors coûts/échéances), matrice d'échéance centralisée,
conducteur FK tenant, pièces liées typées (avec/sans fichier, jamais un coût), historique, notes_internes filtrées
côté API pour read_only, exports read_only GET audités, isolation tenant fail-closed, compat lecture Phase 3.
Aucune donnée Journal réelle ; aucune migration."""
import csv
import io
import re
import uuid
from datetime import datetime, timedelta, timezone

import requests
from dotenv import dotenv_values
from pymongo import MongoClient

_BASE = (dotenv_values("/app/frontend/.env").get("REACT_APP_BACKEND_URL") or "").rstrip("/")
_ENV = dotenv_values("/app/backend/.env")
_RUN = uuid.uuid4().hex[:8]
TENANT_A = f"pytest-fin-a-{_RUN}"
TENANT_B = f"pytest-fin-b-{_RUN}"
ADMIN_A = (f"fin-adm-a-{_RUN}@pytest.ch", f"FinAdmA-{_RUN}-1")
RO_A = (f"fin-ro-a-{_RUN}@pytest.ch", f"FinRoA-{_RUN}-1")
ADMIN_B = (f"fin-adm-b-{_RUN}@pytest.ch", f"FinAdmB-{_RUN}-1")
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 120
STATUSES = ["recue", "a_analyser", "conducteur_a_identifier", "en_attente_conducteur", "contestee",
            "a_payer", "payee", "refacturee", "cloturee", "annulee"]
JOURNAL = {"received": "recue", "to_analyze": "a_analyser", "driver_to_identify": "conducteur_a_identifier",
           "awaiting_driver": "en_attente_conducteur", "disputed": "contestee", "to_pay": "a_payer",
           "paid": "payee", "recharged": "refacturee", "closed": "cloturee", "cancelled": "annulee"}
INFRACTION_TYPES = ["speeding", "parking", "red_light", "toll", "forbidden_zone", "phone", "seatbelt", "other"]  # enum Journal prouvée
INFRACTION_LABELS_FR = {"speeding": "Excès de vitesse", "parking": "Stationnement", "red_light": "Feu rouge", "toll": "Péage",
                        "forbidden_zone": "Zone interdite", "phone": "Téléphone au volant", "seatbelt": "Ceinture de sécurité", "other": "Autre"}
DEADLINE_INACTIVE = {"payee", "refacturee", "cloturee", "annulee"}
_D = lambda n: (datetime.now(timezone.utc) + timedelta(days=n)).strftime("%Y-%m-%d")  # noqa: E731
AMENDE = {"autorite": "Police cantonale vaudoise", "date_infraction": _D(-10), "montant": 120.0, "devise": "CHF",
          "delai_paiement": _D(30), "motif": "Courrier papier non scanné (pytest Lot D)"}
_S, _cache = {}, {}


def _mongo():
    return MongoClient(_ENV["MONGO_URL"])[_ENV["DB_NAME"]]


def _h(creds=ADMIN_A):
    if creds not in _cache:
        r = requests.post(f"{_BASE}/api/auth/login", json={"email": creds[0], "password": creds[1]}, timeout=30)
        assert r.status_code == 200, r.text
        _cache[creds] = {"Authorization": f"Bearer {r.json()['token']}"}
    return _cache[creds]


def sa():
    return _h((_ENV["SUPERADMIN_EMAIL"], _ENV["SUPERADMIN_PASSWORD"]))


def _req(method, path, body=None, creds=ADMIN_A, **kw):
    return requests.request(method, f"{_BASE}/api{path}", json=body, headers=_h(creds), timeout=60, **kw)


def _fine(vehicle_id=None, creds=ADMIN_A, **over):
    body = {**AMENDE, "numero_amende": f"VD-{_RUN}-{uuid.uuid4().hex[:5]}", **over}
    return _req("POST", f"/vehicles/{vehicle_id or _S['veh_a']}/fines", body, creds)


def _mk(**over):
    r = _fine(**over)
    assert r.status_code == 200, r.text
    return r.json()["document_id"]


def _get(doc_id, creds=ADMIN_A):
    r = _req("GET", f"/fines/{doc_id}", creds=creds)
    assert r.status_code == 200, r.text
    return r.json()


def _status(doc_id, status, creds=ADMIN_A, **extra):
    return _req("POST", f"/documents/{doc_id}/fine-status", {"fine_status": status, **extra}, creds)


def _paid(doc_id, payee=True, creds=ADMIN_A, **extra):
    return _req("POST", f"/documents/{doc_id}/paid", {"payee": payee, **extra}, creds)


def _db(doc_id):
    return _mongo().documents.find_one({"id": doc_id}, {"_id": 0})


def _audits(doc_id, action=None, tenant=TENANT_A):
    q = {"tenant_id": tenant, "entity": "document", "entity_id": doc_id}
    if action:
        q["action"] = action
    return list(_mongo().audit_logs.find(q, {"_id": 0}).sort("created_at", 1))


def _cost_ids(creds=ADMIN_A):
    r = _req("GET", "/costs", creds=creds)
    assert r.status_code == 200
    return {i["document_id"] for i in r.json()["items"] if i.get("document_id")}


def _dl_ids(creds=ADMIN_A):
    r = _req("GET", "/deadlines", creds=creds)
    assert r.status_code == 200
    return {i["document_id"] for i in r.json()["items"] if i.get("document_id")}


def _vehicle(payload, creds=ADMIN_A):
    r = _req("POST", "/vehicles", payload, creds)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def setup_module():
    for tenant, admin in ((TENANT_A, ADMIN_A), (TENANT_B, ADMIN_B)):
        assert requests.post(f"{_BASE}/api/admin/tenants", json={"name": tenant, "id": tenant}, headers=sa(), timeout=30).status_code == 200
        assert requests.post(f"{_BASE}/api/admin/tenants/{tenant}/users", json={"email": admin[0], "password": admin[1], "role": "admin"},
                             headers=sa(), timeout=30).status_code == 200
    assert requests.post(f"{_BASE}/api/admin/tenants/{TENANT_A}/users", json={"email": RO_A[0], "password": RO_A[1], "role": "read_only"},
                         headers=sa(), timeout=30).status_code == 200
    _S["veh_a"] = _vehicle({"plaque": f"VD {_RUN[:6].upper()}", "kilometrage": 1000})
    _S["veh_b"] = _vehicle({"plaque": f"BE {_RUN[:6].upper()}"}, creds=ADMIN_B)
    r = _req("POST", "/drivers", {"nom": "Favre", "prenom": "Nina", "matricule_interne": f"F-{_RUN[:4]}"})
    assert r.status_code == 200, r.text
    _S["drv_a"] = r.json()["id"]
    r = _req("POST", "/drivers", {"nom": "Bieri", "prenom": "Tom"}, creds=ADMIN_B)
    assert r.status_code == 200, r.text
    _S["drv_b"] = r.json()["id"]


def teardown_module():
    db = _mongo()
    for tenant in (TENANT_A, TENANT_B):
        for coll in ("users", "vehicles", "documents", "files", "audit_logs", "alerts", "vehicle_field_meta", "tenant_integrations",
                     "doc_categories", "doc_requirements", "tenant_settings", "fuel_transactions", "drivers", "driver_assignments"):
            db[coll].delete_many({"tenant_id": tenant})
        db.tenants.delete_many({"id": tenant})


class TestStatutsParite:
    def test_01_dix_statuts_exacts_exposes(self):
        r = _req("GET", "/fines")
        assert r.status_code == 200, r.text
        assert [s["code"] for s in r.json()["statuses"]] == STATUSES
        assert all(s["label"] for s in r.json()["statuses"])
        assert [t["code"] for t in r.json()["infraction_types"]] == INFRACTION_TYPES  # 8 codes Journal exacts, ordre source
        assert {t["code"]: t["label"] for t in r.json()["infraction_types"]} == INFRACTION_LABELS_FR
        assert "documents" not in INFRACTION_TYPES and "forbidden_zone" in INFRACTION_TYPES
        assert [p["code"] for p in r.json()["piece_types"]] == ["pdf", "photo", "courrier", "contestation", "preuve_paiement", "libre"]

    def test_02_creation_defaut_a_payer_non_payee(self):
        doc_id = _mk()
        _S["doc"] = doc_id
        d = _get(doc_id)
        assert d["fine_status"] == "a_payer" and d["payee"] is False and d["paid_on"] is None and d["paid_at"] is None
        assert d["payment_ref"] is None and d["deadline_active"] is True and d["cost_counted"] is True
        assert d["type_infraction"] == "other" and d["statut"] == "A_PAYER" and d["attachments_count"] == 0
        assert doc_id in _cost_ids() and doc_id in _dl_ids()

    def test_03_mapping_journal_1_pour_1_sans_perte(self):
        for j, code in JOURNAL.items():
            if code in DEADLINE_INACTIVE:
                continue  # états terminaux : jamais à la création (test_03b), mapping couvert par les transitions
            doc_id = _mk(fine_status=j, delai_paiement=None, motif="test mapping")
            assert _get(doc_id)["fine_status"] == code, (j, code)
        r = _fine(fine_status="inconnu_xyz")
        assert r.status_code == 422 and "fine_status inconnu" in r.text
        r = _status(_S["doc"], "to_pay")  # code Journal accepté en transition → code Documents
        assert r.status_code == 200 and r.json()["fine_status"] == "a_payer"
        assert _status(_S["doc"], "n_importe_quoi").status_code == 422
        for j, code in (("paid", "payee"), ("recharged", "refacturee"), ("closed", "cloturee"), ("cancelled", "annulee")):
            tmp = _mk()
            r = _status(tmp, j, motif="mapping terminal", **({"paid_on": _D(-1)} if code in ("payee", "refacturee") else {}))
            assert r.status_code == 200 and r.json()["fine_status"] == code, (j, code)

    def test_03b_creation_directe_etat_terminal_ou_paiement_interdite(self):
        n = _mongo().documents.count_documents({"tenant_id": TENANT_A})
        for code in ("payee", "refacturee", "cloturee", "annulee", "paid", "recharged", "closed", "cancelled"):
            r = _fine(fine_status=code)
            assert r.status_code == 422 and "création directe" in r.text, (code, r.text)
        assert _fine(fine_status="annulee", source="legacy_import", legacy_source="j-test", legacy_id=f"X-{_RUN}", motif=None).status_code == 422
        r = _fine(paid_on=_D(-1))
        assert r.status_code == 422 and "paid_on / payment_ref interdits" in r.text
        assert _fine(payment_ref="REF").status_code == 422
        assert _mongo().documents.count_documents({"tenant_id": TENANT_A}) == n  # aucune écriture
        assert _mongo().documents.count_documents({"tenant_id": TENANT_A, "legacy_id": f"X-{_RUN}"}) == 0

    def test_04_compat_lecture_phase3_sans_fine_status_aucune_ecriture(self):
        now = datetime.now(timezone.utc).isoformat()
        base = {"tenant_id": TENANT_A, "vehicle_id": _S["veh_a"], "folder": "Amendes", "document_type": "amende",
                "business_category": "AMENDE", "extraction_status": "validated", "is_deleted": False, "montant": 60.0,
                "devise": "CHF", "fournisseur": "Police P3", "date_debut": "2026-05-01", "date_expiration": "2026-05-31",
                "created_at": now, "storage_path": None, "justificatif_absent": True, "pages": []}
        unpaid, paid = str(uuid.uuid4()), str(uuid.uuid4())
        _mongo().documents.insert_many([{**base, "id": unpaid, "numero": f"P3-U-{_RUN}", "payee": False, "paid_at": None},
                                        {**base, "id": paid, "numero": f"P3-P-{_RUN}", "payee": True, "paid_at": now}])
        u, p = _get(unpaid), _get(paid)
        assert u["fine_status"] == "a_payer" and u["payee"] is False and u["deadline_active"] is True
        assert p["fine_status"] == "payee" and p["payee"] is True and p["deadline_active"] is False and p["paid_on"] is None
        assert "fine_status" not in _db(unpaid) and "fine_status" not in _db(paid)  # lecture seule : jamais reclassé en base
        assert unpaid in _dl_ids() and paid not in _dl_ids() and unpaid in _cost_ids() and paid in _cost_ids()
        _S["p3_unpaid"], _S["p3_paid"] = unpaid, paid


class TestPaiement:
    def test_05_marquer_payee_paid_on_metier_paid_at_technique(self):
        doc_id = _mk()
        _S["pay"] = doc_id
        r = _paid(doc_id, True, paid_on="2026-06-20", payment_ref="  QR-7781  ")
        assert r.status_code == 200, r.text
        d = _db(doc_id)
        assert d["fine_status"] == "payee" and d["payee"] is True and d["paid_on"] == "2026-06-20"
        assert d["paid_at"] and re.match(r"^\d{4}-\d{2}-\d{2}T", d["paid_at"]) and d["paid_at"][:10] != "2026-06-20"
        assert d["paid_by"] == ADMIN_A[0] and d["payment_ref"] == "QR-7781"
        api = _get(doc_id)
        assert api["payee"] is True and api["deadline_active"] is False and api["cost_counted"] is True and api["statut"] == "PAYEE"
        assert doc_id not in _dl_ids() and doc_id in _cost_ids()
        a = _audits(doc_id, "fine_paid")
        assert len(a) == 1 and "paid_on (date métier) : — → 2026-06-20" in a[0]["detail"] and "payment_ref : — → QR-7781" in a[0]["detail"]
        assert "paid_at (technique) : — → " in a[0]["detail"] and a[0]["user"] == ADMIN_A[0]

    def test_06_payment_ref_optionnelle_jamais_inventee(self):
        doc_id = _mk()
        assert _paid(doc_id, True, paid_on="2026-06-21").status_code == 200
        d = _db(doc_id)
        assert d["fine_status"] == "payee" and d["payment_ref"] is None and d["paid_on"] == "2026-06-21"
        assert _paid(doc_id, True, paid_on="21.06.2026").status_code == 422
        assert _paid(doc_id, True, payment_ref="x" * 121).status_code == 422
        r = _req("PATCH", f"/documents/{doc_id}", {"payment_ref": "VIR-2026-001"})
        assert r.status_code == 200 and _db(doc_id)["payment_ref"] == "VIR-2026-001"
        a = _audits(doc_id, "fine_payment_ref")
        assert len(a) == 1 and "— → VIR-2026-001" in a[0]["detail"]
        _S["pay2"] = doc_id

    def test_07_de_paiement_motif_obligatoire_audit_explicite(self):
        doc_id = _S["pay"]
        before = _db(doc_id)
        r = _paid(doc_id, False)
        assert r.status_code == 422 and "motif obligatoire" in r.text
        assert _db(doc_id)["fine_status"] == "payee"  # rien n'a bougé sans motif
        r = _paid(doc_id, False, motif="Paiement saisi sur la mauvaise amende")
        assert r.status_code == 200, r.text
        d = _db(doc_id)
        assert d["fine_status"] == "a_payer" and d["payee"] is False
        assert d["paid_on"] is None and d["paid_at"] is None and d["payment_ref"] is None and d["paid_by"] is None
        a = _audits(doc_id, "fine_payment_reverted")
        assert len(a) == 1 and a[0]["user"] == ADMIN_A[0] and a[0]["created_at"]
        det = a[0]["detail"]
        assert "motif : Paiement saisi sur la mauvaise amende" in det and "statut payee → a_payer" in det
        assert f"paid_on={before['paid_on']}" in det and f"paid_at={before['paid_at']}" in det and "payment_ref=QR-7781" in det
        assert f"paid_by={ADMIN_A[0]}" in det and "après : paid_on=—, paid_at=—, payment_ref=—" in det
        assert not _audits(doc_id, "fine_unpaid")  # plus d'audit générique
        assert doc_id in _dl_ids()
        assert _paid(doc_id, False, motif="déjà non payée").status_code == 422  # rien à annuler

    def test_08_transition_quitte_payee_motif_et_audit(self):
        doc_id = _S["pay2"]
        assert _db(doc_id)["fine_status"] == "payee"
        assert _status(doc_id, "contestee").status_code == 422
        r = _status(doc_id, "contestee", motif="Contestation déposée après paiement")
        assert r.status_code == 200, r.text
        d = _db(doc_id)
        assert d["fine_status"] == "contestee" and d["payee"] is False and d["paid_on"] is None and d["payment_ref"] is None
        a = _audits(doc_id, "fine_payment_reverted")
        assert len(a) == 1 and "payment_ref=VIR-2026-001" in a[0]["detail"] and "paid_on=2026-06-21" in a[0]["detail"]
        assert not _audits(doc_id, "fine_dispute")

    def test_09_paid_on_jamais_derive_et_refuse_hors_statut_paye(self):
        doc_id = _mk()
        assert _status(doc_id, "a_analyser", paid_on="2026-06-01").status_code == 422
        r = _status(doc_id, "payee")  # sans paid_on : paid_at technique posé, paid_on reste null (jamais substitué)
        assert r.status_code == 200
        d = _db(doc_id)
        assert d["paid_at"] and d["paid_on"] is None and d["payee"] is True
        r = _status(doc_id, "refacturee", paid_on="2026-06-02", payment_ref="REFACT-1")
        assert r.status_code == 200 and r.json()["fine_status"] == "refacturee" and r.json()["payee"] is True
        d = _db(doc_id)
        assert d["paid_on"] == "2026-06-02" and d["payment_ref"] == "REFACT-1" and d["paid_at"]
        assert _audits(doc_id, "fine_recharge")
        _S["refact"] = doc_id

    def test_10_cloture_conserve_les_faits_de_paiement(self):
        doc_id = _S["refact"]
        r = _status(doc_id, "cloturee", motif="Dossier clos")
        assert r.status_code == 200 and r.json()["fine_status"] == "cloturee" and r.json()["payee"] is True
        d = _db(doc_id)
        assert d["paid_on"] == "2026-06-02" and d["payment_ref"] == "REFACT-1" and d["paid_at"]
        assert not _audits(doc_id, "fine_payment_reverted") and _audits(doc_id, "fine_close")
        other = _mk()
        assert _status(other, "cloturee").status_code == 200 and _get(other)["payee"] is False  # clôture sans paiement = non payée
        assert other in _cost_ids() and other not in _dl_ids()


class TestD9EtMatrices:
    def test_11_annulee_motif_obligatoire_hors_couts_hors_echeances_montant_conserve(self):
        doc_id = _mk(montant=250.0)
        assert _status(doc_id, "annulee").status_code == 422
        assert _status(doc_id, "annulee", motif="ab").status_code == 422
        r = _status(doc_id, "annulee", motif="Amende retirée par l'autorité")
        assert r.status_code == 200, r.text
        d = _db(doc_id)
        assert d["fine_status"] == "annulee" and d["montant"] == 250.0 and d["is_deleted"] is False
        assert d["cancel_motif"] == "Amende retirée par l'autorité" and d["cancelled_by"] == ADMIN_A[0] and d["cancelled_at"]
        api = _get(doc_id)
        assert api["cost_counted"] is False and api["deadline_active"] is False and api["statut"] == "ANNULEE" and api["montant"] == 250.0
        assert doc_id not in _cost_ids() and doc_id not in _dl_ids()
        a = _audits(doc_id, "fine_cancel")
        assert len(a) == 1 and "a_payer" in a[0]["detail"] and "annulee" in a[0]["detail"] and "montant conservé : 250.0 CHF" in a[0]["detail"]
        assert "exclue des coûts et des échéances actives (D9)" in a[0]["detail"]
        _S["cancel"] = doc_id

    def test_12_matrice_echeance_et_couts_centralisee_pour_les_10_statuts(self):
        ids = {}
        for s in STATUSES:
            ids[s] = _mk()
            if s != "a_payer":
                assert _status(ids[s], s, motif="matrice pytest").status_code == 200, s
        dl, costs = _dl_ids(), _cost_ids()
        for s, doc_id in ids.items():
            api = _get(doc_id)
            assert api["deadline_active"] is (s not in DEADLINE_INACTIVE), s
            assert (doc_id in dl) is (s not in DEADLINE_INACTIVE), s  # /api/deadlines applique la MÊME règle (fonction unique)
            assert api["cost_counted"] is (s != "annulee"), s
            assert (doc_id in costs) is (s != "annulee"), s  # /api/costs applique la MÊME règle
        rows = {d["id"]: d for d in _req("GET", "/fines", params={"limit": 200}).json()["items"]}
        for s, doc_id in ids.items():
            assert rows[doc_id]["fine_status"] == s
        _S["matrix"] = ids

    def test_13_stats_memes_regles_que_la_liste(self):
        st = _req("GET", "/fines/stats").json()
        lst = _req("GET", "/fines", params={"limit": 200}).json()
        assert st["counts"]["total"] == lst["total"] == sum(st["counts"]["by_status"].values())
        assert st["counts"]["by_status"] == lst["totals"]["by_status"]
        assert st["montants"]["total_compte_chf"] == lst["totals"]["total_chf"]
        assert st["montants"]["annule_chf"] == lst["totals"]["annule_chf"] > 0
        assert st["counts"]["ouvertes"] == sum(1 for d in lst["items"] if d["deadline_active"])
        assert st["counts"]["annulees"] == st["counts"]["by_status"]["annulee"] >= 2
        assert [s["code"] for s in st["statuses"]] == STATUSES and st["urgent_days"] > 0
        assert st["by_type"] and st["top_vehicles"][0]["vehicle_id"] == _S["veh_a"]


class TestConducteur:
    def test_14_driver_fk_tenant_scopee_non_identifie_cross_tenant_404(self):
        doc_id = _S["doc"]
        r = _req("PATCH", f"/documents/{doc_id}", {"driver_id": _S["drv_a"]})
        assert r.status_code == 200 and r.json()["driver_id"] == _S["drv_a"]
        assert _get(doc_id)["driver_nom"].startswith("Nina Favre")
        a = _audits(doc_id, "fine_driver")
        assert len(a) == 1 and "non identifié → Nina Favre" in a[0]["detail"]
        assert _req("PATCH", f"/documents/{doc_id}", {"driver_id": _S["drv_b"]}).status_code == 404  # conducteur d'un autre tenant
        assert _db(doc_id)["driver_id"] == _S["drv_a"]
        r = _req("PATCH", f"/documents/{doc_id}", {"driver_id": ""})
        assert r.status_code == 200 and r.json()["driver_id"] is None and _get(doc_id)["driver_nom"] is None
        assert _fine(driver_id=_S["drv_b"]).status_code == 404
        r = _req("GET", "/fines", params={"driver_id": _S["drv_a"]})
        assert r.status_code == 200 and all(d["driver_id"] == _S["drv_a"] for d in r.json()["items"])


class TestPiecesLiees:
    def test_15_piece_sans_fichier_puis_fichier_meme_piece_jamais_un_cout(self):
        doc_id = _S["doc"]
        n_costs, n_dl, n_fines = len(_cost_ids()), len(_dl_ids()), _req("GET", "/fines").json()["total"]
        r = requests.post(f"{_BASE}/api/documents/{doc_id}/attachments", data={"piece_type": "courrier", "titre": "Courrier autorité 12.06",
                          "date_piece": "2026-06-12", "note": "Original au classeur"}, headers=_h(), timeout=30)
        assert r.status_code == 200, r.text
        att = r.json()
        _S["att"] = att["id"]
        assert att["file_missing"] is True and att["storage_path"] is None and att["parent_document_id"] == doc_id
        assert att["piece_type"] == "courrier" and att["piece_type_label"] == "Courrier" and att["label"] == "Courrier autorité 12.06"
        assert att["tenant_id"] == TENANT_A and att["vehicle_id"] == _S["veh_a"] and "montant" not in att
        assert len(_cost_ids()) == n_costs and len(_dl_ids()) == n_dl and _req("GET", "/fines").json()["total"] == n_fines
        assert _get(doc_id)["attachments_count"] == 1
        assert att["id"] not in {d["id"] for d in _req("GET", "/documents").json()}
        assert att["id"] not in {d["id"] for d in _req("GET", f"/vehicles/{_S['veh_a']}/documents").json()}
        a = _audits(doc_id, "fine_attachment_add")
        assert len(a) == 1 and "sans fichier (métadonnées seules)" in a[0]["detail"] and "aucun coût, aucune échéance" in a[0]["detail"]
        # fichier ajouté plus tard sur la MÊME pièce
        r = requests.post(f"{_BASE}/api/documents/{att['id']}/attach-file", files={"file": ("courrier.png", PNG, "image/png")}, headers=_h(), timeout=60)
        assert r.status_code == 200, r.text
        lst = _req("GET", f"/documents/{doc_id}/attachments").json()
        assert len(lst) == 1 and lst[0]["id"] == att["id"] and lst[0]["file_missing"] is False and lst[0]["storage_path"]
        assert requests.get(f"{_BASE}/api/files/{lst[0]['storage_path']}", headers=_h(), timeout=30).status_code == 200
        assert _audits(doc_id, "fine_attachment_file") and len(_cost_ids()) == n_costs and _get(doc_id)["attachments_count"] == 1

    def test_16_piece_avec_fichier_validations_et_soft_delete(self):
        doc_id = _S["doc"]
        r = requests.post(f"{_BASE}/api/documents/{doc_id}/attachments", data={"piece_type": "preuve_paiement", "titre": "Quittance e-banking"},
                          files={"file": ("quittance.png", PNG, "image/png")}, headers=_h(), timeout=60)
        assert r.status_code == 200, r.text
        att2 = r.json()
        assert att2["file_missing"] is False and att2["storage_path"] and att2["sha256"]
        assert requests.post(f"{_BASE}/api/documents/{doc_id}/attachments", data={"piece_type": "video", "titre": "x y"}, headers=_h(), timeout=30).status_code == 422
        assert requests.post(f"{_BASE}/api/documents/{doc_id}/attachments", data={"piece_type": "pdf", "titre": " "}, headers=_h(), timeout=30).status_code == 422
        assert requests.post(f"{_BASE}/api/documents/{doc_id}/attachments", data={"piece_type": "pdf", "titre": "ok", "date_piece": "12.06.2026"}, headers=_h(), timeout=30).status_code == 422
        assert _get(doc_id)["attachments_count"] == 2
        r = _req("DELETE", f"/documents/{doc_id}/attachments/{att2['id']}")
        assert r.status_code == 200
        db = _db(att2["id"])
        assert db["is_deleted"] is True and db["storage_path"]  # binaire conservé (pas de DELETE objstore)
        assert [a["id"] for a in _req("GET", f"/documents/{doc_id}/attachments").json()] == [_S["att"]]
        assert _audits(doc_id, "fine_attachment_remove")
        assert _req("DELETE", f"/documents/{doc_id}/attachments/{att2['id']}").status_code == 404
        # une pièce n'est jamais une amende
        assert _req("GET", f"/fines/{_S['att']}").status_code == 404
        assert _status(_S["att"], "payee").status_code == 404


class TestHistoriqueEtNotesInternes:
    def test_17_historique_chronologique_par_entite(self):
        rows = _req("GET", f"/documents/{_S['pay']}/history").json()
        actions = [r["action"] for r in rows]
        assert actions[0] == "create" and "fine_paid" in actions and "fine_payment_reverted" in actions
        assert actions.index("fine_paid") < actions.index("fine_payment_reverted")
        assert all(r["user"] == ADMIN_A[0] for r in rows) and rows == sorted(rows, key=lambda r: r["created_at"])
        assert "ip" in rows[0]
        ro = _req("GET", f"/documents/{_S['pay']}/history", creds=RO_A).json()
        assert len(ro) == len(rows) and "ip" not in ro[0]
        assert _req("GET", f"/documents/{_S['pay']}/history", creds=ADMIN_B).status_code == 404

    def test_18_notes_internes_filtrees_cote_api_pour_read_only(self):
        doc_id = _S["doc"]
        r = _req("PATCH", f"/documents/{doc_id}", {"notes_internes": "SECRET-INTERNE-" + _RUN})
        assert r.status_code == 200 and r.json()["notes_internes"] == "SECRET-INTERNE-" + _RUN
        assert _get(doc_id)["notes_internes"] == "SECRET-INTERNE-" + _RUN
        a = _audits(doc_id, "modify")
        assert all("SECRET-INTERNE" not in x["detail"] for x in a) and any("notes_internes modifiées" in x["detail"] for x in a)
        secret = ("SECRET-INTERNE-" + _RUN).encode()
        for path in (f"/fines/{doc_id}", "/fines", "/documents", f"/vehicles/{_S['veh_a']}/documents"):
            r = _req("GET", path, creds=RO_A)
            assert r.status_code == 200, path
            assert secret not in r.content and b"notes_internes" not in r.content, path
        hist = _req("GET", f"/documents/{doc_id}/history", creds=RO_A)
        assert hist.status_code == 200 and secret not in hist.content  # l'audit ne journalise jamais la valeur
        ro_doc = _get(doc_id, creds=RO_A)
        assert "notes_internes" not in ro_doc and ro_doc["fine_status"] == "a_payer"
        assert _get(doc_id)["notes_internes"] == "SECRET-INTERNE-" + _RUN  # admin voit toujours


class TestRbacExportsIsolation:
    def test_19_read_only_mutations_403_lectures_200(self):
        doc_id = _S["doc"]
        assert _paid(doc_id, True, creds=RO_A, paid_on="2026-06-01").status_code == 403
        assert _status(doc_id, "contestee", creds=RO_A).status_code == 403
        assert _req("PATCH", f"/documents/{doc_id}", {"type_infraction": "parking"}, creds=RO_A).status_code == 403
        assert requests.post(f"{_BASE}/api/documents/{doc_id}/attachments", data={"piece_type": "pdf", "titre": "ro"}, headers=_h(RO_A), timeout=30).status_code == 403
        assert _req("DELETE", f"/documents/{doc_id}/attachments/{_S['att']}", creds=RO_A).status_code == 403
        assert _fine(creds=RO_A).status_code == 403
        assert _db(doc_id)["fine_status"] == "a_payer" and _db(doc_id)["type_infraction"] == "other"
        for path in ("/fines", "/fines/stats", f"/fines/{doc_id}", f"/documents/{doc_id}/attachments"):
            assert _req("GET", path, creds=RO_A).status_code == 200, path

    def test_20_exports_csv_xlsx_pdf_read_only_audites_sans_notes_internes(self):
        n_before = _mongo().audit_logs.count_documents({"tenant_id": TENANT_A, "action": "download", "entity": "report"})
        r = requests.get(f"{_BASE}/api/reports/amendes.csv", params={"fine_status": "a_payer,contestee"}, headers=_h(RO_A), timeout=60)
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv") and r.content.startswith("\ufeff".encode("utf-8"))
        text = r.content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text), delimiter=";"))
        assert rows[0][:6] == ["Référence", "Dossier interne", "Véhicule", "Marque / modèle", "Conducteur", "Statut"]
        assert "Payée le" in rows[0] and "Réf. paiement" in rows[0]
        assert "SECRET-INTERNE" not in text and "notes_internes" not in text
        expected = _req("GET", "/fines", params={"fine_status": "a_payer,contestee", "limit": 200}).json()["total"]
        assert len(rows) - 1 == expected > 0
        assert all(row[5] in ("À payer", "Contestée") for row in rows[1:])
        r = requests.get(f"{_BASE}/api/reports/amendes.xlsx", headers=_h(RO_A), timeout=60)
        assert r.status_code == 200 and r.content[:2] == b"PK"
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Amendes", "Synthèse"] and wb["Amendes"].max_row - 1 == _req("GET", "/fines").json()["total"]
        assert "SECRET-INTERNE" not in " ".join(str(c.value or "") for row in wb["Amendes"].iter_rows() for c in row)
        r = requests.get(f"{_BASE}/api/reports/amendes.pdf", params={"vehicle_id": _S["veh_a"]}, headers=_h(RO_A), timeout=60)
        assert r.status_code == 200 and r.content[:4] == b"%PDF"
        assert requests.get(f"{_BASE}/api/reports/amendes.csv", params={"fine_status": "zzz"}, headers=_h(RO_A), timeout=30).status_code == 422
        audits = list(_mongo().audit_logs.find({"tenant_id": TENANT_A, "action": "download", "entity": "report"}).sort("created_at", 1))
        assert len(audits) == n_before + 3
        assert audits[-3]["entity_id"] == "amendes_csv" and audits[-3]["user"] == RO_A[0] and "rôle read_only" in audits[-3]["detail"]
        assert "fine_status" in audits[-3]["detail"] and "a_payer" in audits[-3]["detail"] and audits[-3]["created_at"]
        assert audits[-2]["entity_id"] == "amendes_xlsx" and audits[-1]["entity_id"] == "amendes_pdf"
        # export = GET seulement, aucune écriture métier
        assert _db(_S["doc"])["fine_status"] == "a_payer"
        # payment_ref exportée uniquement si présente
        r = requests.get(f"{_BASE}/api/reports/amendes.csv", params={"fine_status": "cloturee"}, headers=_h(), timeout=60)
        rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
        col = rows[0].index("Réf. paiement")
        refs = {row[col] for row in rows[1:]}
        assert "REFACT-1" in refs and "" in refs

    def test_21_isolation_tenant_fail_closed(self):
        doc_id = _S["doc"]
        assert _req("GET", f"/fines/{doc_id}", creds=ADMIN_B).status_code == 404
        assert _status(doc_id, "contestee", creds=ADMIN_B).status_code == 404
        assert _paid(doc_id, True, creds=ADMIN_B, paid_on="2026-06-01").status_code == 404
        assert _req("GET", f"/documents/{doc_id}/attachments", creds=ADMIN_B).status_code == 404
        assert requests.post(f"{_BASE}/api/documents/{doc_id}/attachments", data={"piece_type": "pdf", "titre": "intrus"}, headers=_h(ADMIN_B), timeout=30).status_code == 404
        assert _req("DELETE", f"/documents/{doc_id}/attachments/{_S['att']}", creds=ADMIN_B).status_code == 404
        b_fine = _mk(vehicle_id=_S["veh_b"], creds=ADMIN_B, montant=77.0)
        lst_b = _req("GET", "/fines", creds=ADMIN_B).json()
        assert [d["id"] for d in lst_b["items"]] == [b_fine] and lst_b["totals"]["total_chf"] == 77.0
        st_b = _req("GET", "/fines/stats", creds=ADMIN_B).json()
        assert st_b["counts"]["total"] == 1 and st_b["montants"]["total_compte_chf"] == 77.0
        r = requests.get(f"{_BASE}/api/reports/amendes.csv", headers=_h(ADMIN_B), timeout=60)
        rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
        assert len(rows) == 2 and f"BE {_RUN[:6].upper()}" in rows[1]
        assert _db(doc_id)["fine_status"] == "a_payer" and _db(doc_id)["tenant_id"] == TENANT_A


class TestFiltresEtCompat:
    def test_22_filtres_liste_pagination_recherche(self):
        parking = _mk(type_infraction="parking", date_infraction="2026-03-05", delai_paiement="2026-04-05", montant=40.0,
                      dossier_interne=f"DOS-{_RUN}", priorite="high", lieu_infraction={"ville": "Lausanne", "canton": "VD"})
        d = _get(parking)
        assert d["type_infraction"] == "parking" and d["type_infraction_label"] == "Stationnement" and d["lieu_label"] == "Lausanne, VD"
        assert d["priorite"] == "high" and d["dossier_interne"] == f"DOS-{_RUN}"
        g = lambda **p: _req("GET", "/fines", params={"limit": 200, **p}).json()  # noqa: E731
        assert {x["id"] for x in g(type_infraction="parking")["items"]} == {parking}
        assert parking in {x["id"] for x in g(date_from="2026-03-01", date_to="2026-03-31")["items"]}
        assert {x["id"] for x in g(date_from="2026-03-01", date_to="2026-03-31")["items"]} == {parking}
        assert {x["id"] for x in g(due_from="2026-04-01", due_to="2026-04-30")["items"]} == {parking}
        assert {x["id"] for x in g(montant_min=30, montant_max=50)["items"]} == {parking}
        assert {x["id"] for x in g(q=f"dos-{_RUN}")["items"]} == {parking}
        assert {x["id"] for x in g(q="lausanne")["items"]} == {parking}
        assert {x["id"] for x in g(priorite="high")["items"]} == {parking}
        assert {x["id"] for x in g(fine_status="paid,recharged")["items"]} == {x["id"] for x in g(fine_status="payee,refacturee")["items"]}
        assert _req("GET", "/fines", params={"fine_status": "bidon"}).status_code == 422
        assert _req("GET", "/fines", params={"date_from": "01.03.2026"}).status_code == 422
        page = _req("GET", "/fines", params={"limit": 2, "offset": 0, "sort": "-montant"}).json()
        assert len(page["items"]) == 2 and page["limit"] == 2 and page["items"][0]["montant"] >= page["items"][1]["montant"]
        page2 = _req("GET", "/fines", params={"limit": 2, "offset": 2, "sort": "-montant"}).json()
        assert {x["id"] for x in page2["items"]}.isdisjoint({x["id"] for x in page["items"]})
        assert _req("GET", "/fines", params={"vehicle_id": _S["veh_b"]}).json()["total"] == 0  # véhicule d'un autre tenant : rien

    def test_23_type_infraction_enum_8_codes_hors_enum_refuse_historique_conserve(self):
        # les 8 codes acceptés à la création (code technique stocké tel quel, libellé FR exposé) et en PATCH
        created = {}
        for code in INFRACTION_TYPES:
            created[code] = _mk(type_infraction=code, delai_paiement=None)
            d = _get(created[code])
            assert d["type_infraction"] == code and d["type_infraction_label"] == INFRACTION_LABELS_FR[code], code
            assert _db(created[code])["type_infraction"] == code
        r = _req("PATCH", f"/documents/{created['other']}", {"type_infraction": "seatbelt"})
        assert r.status_code == 200 and r.json()["type_infraction"] == "seatbelt"
        assert _get(created["other"])["type_infraction_label"] == "Ceinture de sécurité"
        assert any("type_infraction: other → seatbelt" in a["detail"] for a in _audits(created["other"], "modify"))
        assert {x["id"] for x in _req("GET", "/fines", params={"type_infraction": "forbidden_zone", "limit": 200}).json()["items"]} == {created["forbidden_zone"]}
        by_type = {t["code"]: t for t in _req("GET", "/fines/stats").json()["by_type"]}
        assert by_type["toll"]["label"] == "Péage" and by_type["toll"]["count"] >= 1
        assert _get(_mk(type_infraction=None, delai_paiement=None))["type_infraction"] == "other"  # non fourni → défaut `other`
        # nouvelle saisie hors enum : refusée (création, PATCH, legacy_import, null, vide) — aucune écriture
        n = _mongo().documents.count_documents({"tenant_id": TENANT_A})
        for bad in ("documents", "code_legacy_inconnu", "Speeding", "red light"):
            r = _fine(type_infraction=bad)
            assert r.status_code == 422 and "type_infraction" in r.text, (bad, r.text)
        assert _fine(type_infraction="documents", source="legacy_import", legacy_source="j-test", legacy_id=f"TI-{_RUN}", motif=None).status_code == 422
        assert _req("PATCH", f"/documents/{created['other']}", {"type_infraction": "documents"}).status_code == 422
        assert _req("PATCH", f"/documents/{created['other']}", {"type_infraction": None}).status_code == 422
        assert _req("PATCH", f"/documents/{created['other']}", {"type_infraction": ""}).status_code == 422
        assert _mongo().documents.count_documents({"tenant_id": TENANT_A}) == n and _db(created["other"])["type_infraction"] == "seatbelt"
        # valeur historique déjà stockée hors enum (donnée réelle non modifiée) : lue et affichée telle quelle, jamais convertie en `other`
        hist = str(uuid.uuid4())
        _mongo().documents.insert_one({"tenant_id": TENANT_A, "vehicle_id": _S["veh_a"], "id": hist, "folder": "Amendes", "document_type": "amende",
                                       "business_category": "AMENDE", "extraction_status": "validated", "is_deleted": False, "montant": 30.0,
                                       "devise": "CHF", "fournisseur": "Police P3", "numero": f"P3-T-{_RUN}", "type_infraction": "code_legacy_inconnu",
                                       "payee": False, "created_at": datetime.now(timezone.utc).isoformat(), "pages": []})
        d = _get(hist)
        assert d["type_infraction"] == "code_legacy_inconnu" and d["type_infraction_label"] == "code_legacy_inconnu"
        r = _req("PATCH", f"/documents/{hist}", {"dossier_interne": "DOS-HIST"})  # autre champ : OK, la valeur historique reste intacte
        assert r.status_code == 200 and _db(hist)["type_infraction"] == "code_legacy_inconnu"
        assert _req("GET", "/fines/stats").json()["by_type"] and any(t["code"] == "code_legacy_inconnu" for t in _req("GET", "/fines/stats").json()["by_type"])
        # legacy_import valide (code enum) pour le test d'idempotence suivant
        _S["legacy"] = _mk(type_infraction="red_light", source="legacy_import", legacy_source="journal-test", legacy_id=f"FINE-{_RUN}-X",
                           fine_status="disputed", motif=None)
        assert _get(_S["legacy"])["fine_status"] == "contestee" and _get(_S["legacy"])["type_infraction_label"] == "Feu rouge"

    def test_24_rejeu_legacy_idempotent_ne_reclasse_pas_le_statut(self):
        doc_id = _S["legacy"]
        assert _status(doc_id, "a_payer").status_code == 200
        r = _fine(type_infraction="red_light", source="legacy_import", legacy_source="journal-test",
                  legacy_id=f"FINE-{_RUN}-X", fine_status="to_analyze", motif=None)
        assert r.status_code == 200 and r.json()["created"] is False and r.json()["document_id"] == doc_id
        assert _db(doc_id)["fine_status"] == "a_payer"  # l'état métier courant n'est jamais écrasé par un rejeu
        assert _mongo().documents.count_documents({"tenant_id": TENANT_A, "legacy_id": f"FINE-{_RUN}-X"}) == 1
        assert _audits(doc_id, "legacy_replay")

    def test_25_d7_devise_etrangere_pending_fx_hors_totaux(self):
        eur = _mk(devise="EUR", montant=90.0)
        d = _get(eur)
        assert d["pending_fx"] is True and d["montant_chf_effectif"] is None and d["cost_counted"] is True
        lst = _req("GET", "/fines", params={"limit": 200}).json()
        assert lst["totals"]["pending_fx_count"] >= 1
        st = _req("GET", "/fines/stats").json()
        assert st["montants"]["pending_fx_count"] >= 1
        assert eur not in _cost_ids()  # exclu des items (pending_fx) tant que la contre-valeur manque
        r = _req("PATCH", f"/documents/{eur}", {"montant_chf": 85.5})
        assert r.status_code == 200 and _get(eur)["montant_chf_effectif"] == 85.5 and eur in _cost_ids()

    def test_26_statique_regle_unique_et_aucune_donnee_journal(self):
        src = open("/app/backend/server.py", encoding="utf-8").read()
        fines_src = open("/app/backend/fines.py", encoding="utf-8").read()
        assert src.count("fin.deadline_active(") >= 2 and src.count("fin.cost_counted(") >= 2  # échéances + coûts + sortie API
        assert "def is_paid" in fines_src and fines_src.count("def is_paid") == 1
        assert 'in ("payee", "refacturee", "cloturee", "annulee")' not in src  # pas de matrice recopiée hors fines.py
        body = src[src.index("async def collect_costs"):]
        body = body[:body.find("\nasync def ", 10)]
        assert "fuel_transactions" not in body
        assert "AMD-2026-" not in src and "AMD-2026-" not in fines_src  # aucun identifiant Journal réel
